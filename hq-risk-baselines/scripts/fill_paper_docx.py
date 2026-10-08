"""Paste the sections of notes/paper/sections.md and the compact results table into the IEEE Word draft.
The draft stays outside the repo; this script only writes a copy next to it."""

import argparse
import io
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.build_tables import build_all

# Word saved the draft as Strict OOXML; python-docx reads Transitional. Same markup, different namespace URIs.
STRICT = "http://purl.oclc.org/ooxml/"
TRANSITIONAL = {
    "officeDocument/relationships": "officeDocument/2006/relationships", "officeDocument/math": "officeDocument/2006/math",
    "officeDocument/sharedTypes": "officeDocument/2006/sharedTypes", "officeDocument/customXml": "officeDocument/2006/customXml",
    "officeDocument/extendedProperties": "officeDocument/2006/extended-properties",
    "officeDocument/customProperties": "officeDocument/2006/custom-properties",
    "officeDocument/docPropsVTypes": "officeDocument/2006/docPropsVTypes", "officeDocument/bibliography": "officeDocument/2006/bibliography",
    "wordprocessingml/main": "wordprocessingml/2006/main", "drawingml/main": "drawingml/2006/main",
    "drawingml/wordprocessingDrawing": "drawingml/2006/wordprocessingDrawing", "drawingml/picture": "drawingml/2006/picture",
    "schemaLibrary/main": "schemaLibrary/2006/main",
}
FONT_PT = 8           # Table I in the draft uses 8 pt runs
INDENT_PT = 13.5      # first-line indent of the draft's body paragraphs
COLUMN_PT = [70, 32, 47, 47, 47]  # sums to one IEEE column (≈ 243 pt)
CELL_MARGIN_PT = 2


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--draft", required=True, help="the lead's .docx")
    parser.add_argument("--out", required=True)
    parser.add_argument("--sections", default="notes/paper/sections.md")
    parser.add_argument("--summary", default="results/summary/summary.csv")
    parser.add_argument("--author", required=True, help="name on the Word comments")
    parser.add_argument("--table-under", default="Results", help="heading that receives the compact table")
    return parser.parse_args()


def transitional_copy(path: str) -> io.BytesIO:
    """Rewrite Strict namespace URIs to their Transitional equivalents; returns the draft unchanged if it is not Strict."""
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith((".xml", ".rels")) and STRICT.encode() in data:
                text = data.decode()
                for strict, trans in TRANSITIONAL.items():
                    text = text.replace(STRICT + strict, "http://schemas.openxmlformats.org/" + trans)
                data = text.encode()
            dst.writestr(item, data)
    out.seek(0)
    return out


def read_sections(path: Path) -> dict[str, list[str]]:
    sections = {}
    for chunk in re.split(r"^## ", path.read_text(), flags=re.M)[1:]:
        heading, _, body = chunk.partition("\n")
        sections[heading.strip()] = [p.strip().replace("\n", " ") for p in re.split(r"\n\s*\n", body) if p.strip()]
    return sections


def paragraph_by_text(doc: Document, text: str):
    match = [p for p in doc.paragraphs if p.text.strip().lower() == text.lower()]
    if not match:
        raise LookupError(f"no paragraph reads {text!r}")
    return match[0]


def remove_empty_after(paragraph) -> None:
    nxt = paragraph._p.getnext()
    if nxt is not None and nxt.tag.endswith("}p") and not "".join(nxt.itertext()).strip():
        nxt.getparent().remove(nxt)


def body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Pt(INDENT_PT)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def styled_paragraph(doc: Document, text: str, style: str):
    p = doc.add_paragraph(text)
    p.style = doc.styles[style]
    return p


def set_cell_margins(table, pt: int) -> None:
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(pt * 20))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    table._tbl.tblPr.append(margins)


def set_cell(cell, text: str, bold: bool = False, width_pt: int | None = None) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text.replace(" ± ", "\u00a0±\u00a0"))
    run.font.size = Pt(FONT_PT)
    run.bold = bold
    if width_pt:
        cell.width = Pt(width_pt)


def compact_docx_table(doc: Document, table: dict):
    t = doc.add_table(rows=1 + len(table["rows"]), cols=len(table["columns"]))
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    set_cell_margins(t, CELL_MARGIN_PT)
    for j, width in enumerate(COLUMN_PT):
        t.columns[j].width = Pt(width)
    for j, name in enumerate(table["columns"]):
        set_cell(t.rows[0].cells[j], name, bold=True, width_pt=COLUMN_PT[j])
    for i, row in enumerate(table["rows"], start=1):
        for j, c in enumerate(row):
            set_cell(t.rows[i].cells[j], c["text"], width_pt=COLUMN_PT[j])
    return t


def insert_after(anchor, elements) -> None:
    for el in elements:
        anchor.addnext(el)
        anchor = el


def main() -> None:
    args = parse_args()
    doc = Document(transitional_copy(args.draft))
    sections = read_sections(Path(args.sections))
    compact = build_all(pd.read_csv(args.summary))[-1]

    for heading, paragraphs in sections.items():
        if heading.startswith("Comment:"):
            target = paragraph_by_text(doc, heading.split(":", 1)[1].strip())
            doc.add_comment(target.runs, text=" ".join(paragraphs), author=args.author, initials=args.author[:2].upper())
            continue
        anchor = paragraph_by_text(doc, heading)
        remove_empty_after(anchor)
        new = []
        if heading == args.table_under:
            new.append(styled_paragraph(doc, compact["caption"], "table head")._p)
            new.append(compact_docx_table(doc, compact)._tbl)
            note = styled_paragraph(doc, compact["note"], "table footnote")
            note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            new.append(note._p)
        new += [body_paragraph(doc, p)._p for p in paragraphs]
        insert_after(anchor._p, new)

    doc.save(args.out)
    words = sum(len(p.split()) for ps in sections.values() for p in ps)
    print(f"{args.out}: {len(sections)} sections, {words} words, table with {len(compact['rows'])} rows")


if __name__ == "__main__":
    main()
